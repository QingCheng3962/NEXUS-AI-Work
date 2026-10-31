"""文档读取：Word / PDF / Excel / 文本转成可分析的纯文本。"""

from __future__ import annotations

import re
import uuid
from dataclasses import dataclass
from pathlib import Path

from . import settings

MAX_UPLOAD_BYTES = 25 * 1024 * 1024
MAX_CHARS = 120_000
SUPPORTED_SUFFIXES = {".pdf", ".docx", ".xlsx", ".xlsm", ".txt", ".md", ".csv"}
UPLOAD_DIR = settings.DATA_DIR / "uploads"


class DocumentError(RuntimeError):
    """可展示给用户的文档读取错误。"""


@dataclass
class DocumentText:
    title: str
    source: str
    text: str
    chars: int
    pages: int = 0
    sheets: int = 0
    truncated: bool = False


def _clean(text: str) -> str:
    text = text.replace("\x00", " ")
    text = re.sub(r"\r\n?", "\n", text)
    text = re.sub(r"[ \t]+\n", "\n", text)
    text = re.sub(r"\n{4,}", "\n\n\n", text)
    return text.strip()


def _finish(title: str, source: str, text: str, *, pages: int = 0, sheets: int = 0) -> DocumentText:
    clean = _clean(text)
    truncated = len(clean) > MAX_CHARS
    if truncated:
        clean = clean[:MAX_CHARS] + "\n\n…（文档过长，已截断）"
    if not clean:
        raise DocumentError("文档中没有可提取的文字。")
    return DocumentText(
        title=title,
        source=source,
        text=clean,
        chars=len(clean),
        pages=pages,
        sheets=sheets,
        truncated=truncated,
    )


def _read_pdf(path: Path) -> DocumentText:
    try:
        from pypdf import PdfReader
    except ImportError as exc:  # pragma: no cover - 依赖已声明
        raise DocumentError("缺少 PDF 解析组件，请重新安装依赖。") from exc
    try:
        reader = PdfReader(str(path))
    except Exception as exc:  # noqa: BLE001 - 第三方解析异常统一展示
        raise DocumentError(f"无法读取 PDF：{exc}") from exc
    pages: list[str] = []
    for index, page in enumerate(reader.pages[:300]):
        try:
            pages.append(f"\n--- 第 {index + 1} 页 ---\n{page.extract_text() or ''}")
        except Exception:
            continue
    return _finish(path.stem, path.name, "\n".join(pages), pages=len(reader.pages))


def _read_docx(path: Path) -> DocumentText:
    try:
        from docx import Document
    except ImportError as exc:  # pragma: no cover
        raise DocumentError("缺少 Word 解析组件，请重新安装依赖。") from exc
    try:
        document = Document(str(path))
    except Exception as exc:  # noqa: BLE001
        raise DocumentError(f"无法读取 Word 文档：{exc}") from exc
    parts = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))
    return _finish(path.stem, path.name, "\n".join(parts))


def _read_excel(path: Path) -> DocumentText:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover
        raise DocumentError("缺少 Excel 解析组件，请重新安装依赖。") from exc
    try:
        workbook = load_workbook(str(path), read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001
        raise DocumentError(f"无法读取 Excel：{exc}") from exc
    parts: list[str] = []
    sheet_count = 0
    try:
        for sheet in workbook.worksheets[:20]:
            sheet_count += 1
            parts.append(f"\n--- 工作表：{sheet.title} ---")
            for row_index, row in enumerate(sheet.iter_rows(values_only=True)):
                if row_index >= 2000:
                    parts.append("…（该工作表行数过多，已截断）")
                    break
                values = [str(value).strip() for value in row if value is not None and str(value).strip()]
                if values:
                    parts.append(" | ".join(values))
    finally:
        workbook.close()
    return _finish(path.stem, path.name, "\n".join(parts), sheets=sheet_count)


def _read_text(path: Path) -> DocumentText:
    data = path.read_bytes()
    for encoding in ("utf-8-sig", "utf-8", "gb18030", "utf-16"):
        try:
            text = data.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    else:
        raise DocumentError("文本编码无法识别，请另存为 UTF-8 后重试。")
    return _finish(path.stem, path.name, text)


def extract_text(path: str | Path) -> DocumentText:
    source = Path(path).expanduser().resolve()
    if not source.exists() or not source.is_file():
        raise DocumentError("找不到上传的文档。")
    suffix = source.suffix.lower()
    if suffix == ".pdf":
        return _read_pdf(source)
    if suffix == ".docx":
        return _read_docx(source)
    if suffix in (".xlsx", ".xlsm"):
        return _read_excel(source)
    if suffix in (".txt", ".md", ".csv"):
        return _read_text(source)
    if suffix == ".doc":
        raise DocumentError("暂不支持旧版 .doc，请先另存为 .docx。")
    raise DocumentError("仅支持 PDF、Word（.docx）、Excel（.xlsx）和文本文件。")


def save_upload(name: str, data: bytes) -> Path:
    """把上传内容写入本机临时目录；文件名由应用生成，不信任客户端路径。"""
    suffix = Path(name or "").suffix.lower()
    if suffix not in SUPPORTED_SUFFIXES:
        raise DocumentError("仅支持 PDF、Word（.docx）、Excel（.xlsx）和文本文件。")
    if len(data) > MAX_UPLOAD_BYTES:
        raise DocumentError("文档超过 25 MB，请先拆分或压缩。")
    UPLOAD_DIR.mkdir(parents=True, exist_ok=True)
    target = UPLOAD_DIR / f"{uuid.uuid4().hex}{suffix}"
    target.write_bytes(data)
    return target


def build_summary_prompt(doc: DocumentText, *, max_chars: int = 14_000) -> str:
    body = doc.text[:max_chars]
    truncated = "（内容已截断）" if len(doc.text) > max_chars else ""
    return f"""请阅读下面的文档并生成一份中文摘要笔记。

文档：{doc.source}
标题：{doc.title}
{truncated}

要求：
1. 先给出 2-4 句话的核心摘要；
2. 再用 Markdown 列表提炼关键观点、数据或结论；
3. 如果文档包含行动项或待办，单独列出；
4. 不要编造文档中没有的信息；
5. 只输出摘要正文，不要解释你的工作过程。

--- 文档内容 ---
{body}
"""


def note_markdown(doc: DocumentText, summary: str) -> str:
    meta = f"> 来源：{doc.source} · {doc.chars} 字"
    if doc.pages:
        meta += f" · {doc.pages} 页"
    if doc.sheets:
        meta += f" · {doc.sheets} 个工作表"
    return f"{meta}\n\n## AI 摘要\n\n{summary.strip()}\n"
