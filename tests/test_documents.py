"""文档提取测试：Word / Excel / 文本。"""

from pathlib import Path

from docx import Document
from openpyxl import Workbook

from core import documents


def test_extract_docx(tmp_path: Path) -> None:
    path = tmp_path / "demo.docx"
    document = Document()
    document.add_paragraph("Anvil 文档标题")
    document.add_paragraph("这是第一段正文。")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "指标"
    table.cell(0, 1).text = "数值"
    document.save(path)

    result = documents.extract_text(path)
    assert "Anvil 文档标题" in result.text
    assert "指标 | 数值" in result.text
    assert result.source == "demo.docx"


def test_extract_xlsx(tmp_path: Path) -> None:
    path = tmp_path / "demo.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "数据"
    sheet.append(["项目", "金额"])
    sheet.append(["服务器", 1200])
    workbook.save(path)

    result = documents.extract_text(path)
    assert "工作表：数据" in result.text
    assert "服务器 | 1200" in result.text
    assert result.sheets == 1


def test_extract_text_file(tmp_path: Path) -> None:
    path = tmp_path / "demo.md"
    path.write_text("# 标题\n\n正文", encoding="utf-8")
    result = documents.extract_text(path)
    assert result.title == "demo"
    assert "正文" in result.text


def test_save_upload_rejects_unknown_suffix(tmp_path: Path, monkeypatch) -> None:
    monkeypatch.setattr(documents, "UPLOAD_DIR", tmp_path / "uploads")
    try:
        documents.save_upload("evil.exe", b"x")
    except documents.DocumentError as exc:
        assert "仅支持" in str(exc)
    else:
        raise AssertionError("should reject unknown suffix")
